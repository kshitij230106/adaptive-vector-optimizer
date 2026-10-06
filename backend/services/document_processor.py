import os
import pymupdf

from docx import Document
from pptx import Presentation

# ============================================================
# PDF
# ============================================================


def extract_text_from_pdf(file_path):

    text = ""

    document = pymupdf.open(file_path)

    try:

        for page in document:

            text += page.get_text()

    finally:

        document.close()

    return text


# ============================================================
# DOCX
# ============================================================


def extract_text_from_docx(file_path):

    document = Document(file_path)

    text_parts = []

    # Paragraphs

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text_parts.append(paragraph.text)

    # Tables

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                if cell.text.strip():

                    row_text.append(cell.text.strip())

            if row_text:

                text_parts.append(" | ".join(row_text))

    return "\n".join(text_parts)


# ============================================================
# PPTX
# ============================================================


def extract_text_from_pptx(file_path):

    presentation = Presentation(file_path)

    text_parts = []

    for slide_number, slide in enumerate(presentation.slides, start=1):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                if shape.text.strip():

                    slide_text.append(shape.text.strip())

        if slide_text:

            text_parts.append(f"Slide {slide_number}\n" + "\n".join(slide_text))

    return "\n\n".join(text_parts)


# ============================================================
# TXT / MD
# ============================================================


def extract_text_from_text_file(file_path):

    with open(file_path, "r", encoding="utf-8", errors="ignore") as file:

        return file.read()


# ============================================================
# MAIN EXTRACTION FUNCTION
# ============================================================


def extract_text(file_path):

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":

        return extract_text_from_pdf(file_path)

    elif extension == ".docx":

        return extract_text_from_docx(file_path)

    elif extension == ".pptx":

        return extract_text_from_pptx(file_path)

    elif extension in [".txt", ".md"]:

        return extract_text_from_text_file(file_path)

    elif extension == ".doc":

        raise Exception(
            "Old .doc files are not supported directly. "
            "Please convert the file to .docx and upload it again."
        )

    else:

        raise Exception(f"Unsupported file type: {extension}")
